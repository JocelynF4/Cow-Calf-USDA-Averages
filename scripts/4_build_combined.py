import csv
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
OUT = os.path.join(ROOT, "Rancher_Lookup_Combined.xlsx")

D = json.load(open(os.path.join(DATA, "combine_data.json")))
hay, fence, margin = D["hay"], D["fence"], D["margin"]

CAT_LABELS = [
    "Overall Feeder", "Feeder Steers", "Feeder Heifers", "Feeder Bulls",
    "Feeder Cows", "Stock Cows ($/cwt)", "Stock Cows ($/head)",
    "Bred Cows ($/head)", "Bred Cows ($/cwt)", "Bred Heifers ($/head)",
    "Bred Heifers ($/cwt)", "Cow-Calf Pairs ($/head)", "Heifer Pairs ($/head)",
    "Replacement Bulls ($/head)", "Replacement Bulls ($/cwt)",
]

cattle = {}
for r in csv.DictReader(open(os.path.join(DATA, "averages_long.csv"))):
    cattle.setdefault(r["state"], {})[r["category"]] = float(r["avg_price"])
cattle_states = sorted(cattle)

counties = []
for r in csv.DictReader(open(os.path.join(DATA, "counties.csv"))):
    counties.append((r["label"], r["county"], r["state_abbr"], r["state_name"],
                     r["fips"], r["ers_region"]))
all_states = sorted({c[3] for c in counties})

NAVY, BLUE, LIGHT, AMBER, GREY, GREEN = "1F3A5F", "2F5C8F", "DCE6F1", "FFF2CC", "F2F2F2", "E2EFDA"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR = Font(bold=True, color="FFFFFF", size=10)
SUB = Font(italic=True, size=10, color="595959")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def style_header(cell, color=BLUE):
    cell.font, cell.alignment, cell.border = HDR, CEN, BORDER
    cell.fill = PatternFill("solid", fgColor=color)


wb = Workbook()
lk = wb.active
lk.title = "Rancher Lookup"

cs = wb.create_sheet("Cattle_by_State")
cs.cell(1, 1, "State")
for j, lab in enumerate(CAT_LABELS, start=2):
    cs.cell(1, j, lab)
for i, st in enumerate(cattle_states, start=2):
    cs.cell(i, 1, st)
    for j, lab in enumerate(CAT_LABELS, start=2):
        v = cattle[st].get(lab)
        if v is not None:
            cs.cell(i, j, v)
cattle_last = 1 + len(cattle_states)
for col in range(1, 2 + len(CAT_LABELS)):
    style_header(cs.cell(1, col))
cs.column_dimensions["A"].width = 16

hs = wb.create_sheet("Hay_by_State")
hs.cell(1, 1, "State")
hs.cell(1, 2, "Avg Hay ($/ton)")
for i, (st, price) in enumerate(sorted(hay.items()), start=2):
    hs.cell(i, 1, st)
    hs.cell(i, 2, price)
hay_last = 1 + len(hay)
for col in (1, 2):
    style_header(hs.cell(1, col))
hs.column_dimensions["A"].width = 18
hs.column_dimensions["B"].width = 16

ms = wb.create_sheet("CalfMargin_by_Region")
ms.cell(1, 1, "ERS Region")
ms.cell(1, 2, "Calf Margin ($/cow)")
for i, (reg, val) in enumerate(sorted(margin.items()), start=2):
    ms.cell(i, 1, reg)
    ms.cell(i, 2, val)
margin_last = 1 + len(margin)
for col in (1, 2):
    style_header(ms.cell(1, col))
ms.column_dimensions["A"].width = 22
ms.column_dimensions["B"].width = 18

fr = wb.create_sheet("Fencing_Rates")
fr.cell(1, 1, "Fence Type")
fr.cell(1, 2, "Cost per Foot")
for i, (t, rate) in enumerate(fence.items(), start=2):
    fr.cell(i, 1, t)
    fr.cell(i, 2, rate)
fence_last = 1 + len(fence)
for col in (1, 2):
    style_header(fr.cell(1, col))
fr.column_dimensions["A"].width = 30
fr.column_dimensions["B"].width = 14

co = wb.create_sheet("Counties")
for j, h in enumerate(["Label", "County", "State", "State Name", "FIPS", "ERS Region"], start=1):
    style_header(co.cell(1, j))
    co.cell(1, j, h)
for i, rec in enumerate(counties, start=2):
    for j, v in enumerate(rec, start=1):
        co.cell(i, j, v)
co_last = 1 + len(counties)
co.column_dimensions["A"].width = 26
co.column_dimensions["D"].width = 16
co.freeze_panes = "A2"

lists = wb.create_sheet("Lists")
lists.cell(1, 1, "States")
for i, st in enumerate(all_states, start=2):
    lists.cell(i, 1, st)
states_last = 1 + len(all_states)
lists.cell(1, 3, "Counties (filtered by selected state)")
lists["C2"] = (f"=IFERROR(SORT(FILTER(Counties!$A$2:$A${co_last},"
               f"Counties!$D$2:$D${co_last}='Rancher Lookup'!$C$6)),\"— pick a state —\")")
lists.column_dimensions["A"].width = 22
lists.column_dimensions["C"].width = 30

wb.defined_names.add(DefinedName("StateList", attr_text=f"Lists!$A$2:$A${states_last}"))
wb.defined_names.add(DefinedName("CountyList", attr_text="Lists!$C$2#"))

lk.sheet_view.showGridLines = False
lk["B2"] = "Rancher County Lookup"
lk["B2"].font = Font(bold=True, size=20, color=NAVY)
lk["B3"] = ("Pick a state, then a county. Auction cattle prices and hay are by state; "
            "calf margin is by the county's ERS region. Fencing is estimated separately below.")
lk["B3"].font = SUB
lk.merge_cells("B3:F3")


def section(cell, text):
    c = lk[cell]
    c.value, c.font = text, Font(bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = LEFT


def field_label(cell, text):
    c = lk[cell]
    c.value, c.font = text, Font(bold=True, size=11, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment, c.border = RIGHT, BORDER


def input_cell(cell, value=None):
    c = lk[cell]
    if value is not None:
        c.value = value
    c.fill = PatternFill("solid", fgColor=AMBER)
    c.alignment, c.border, c.font = LEFT, BORDER, Font(bold=True, size=11)


def out_label(cell, text):
    c = lk[cell]
    c.value, c.font, c.alignment, c.border = text, Font(size=11), LEFT, BORDER
    c.fill = PatternFill("solid", fgColor=GREY)


def out_value(cell, formula, fmt=None):
    c = lk[cell]
    c.value = formula
    c.font, c.alignment, c.border = Font(bold=True, size=11, color=NAVY), CEN, BORDER
    c.fill = PatternFill("solid", fgColor=LIGHT)
    if fmt:
        c.number_format = fmt


section("B5", "  1 · LOCATION")
field_label("B6", "State:")
input_cell("C6", all_states[0] if all_states else "")
field_label("B7", "County:")
input_cell("C7")
lk["B8"].value = "ERS Region:"
lk["B8"].font, lk["B8"].alignment = Font(italic=True, size=10, color="595959"), RIGHT
lk["C8"] = f'=IFERROR(XLOOKUP($C$7,Counties!$A$2:$A${co_last},Counties!$F$2:$F${co_last}),"")'
lk["C8"].font, lk["C8"].alignment = Font(italic=True, size=10, color="595959"), CEN

CATTLE_RANGE = f"Cattle_by_State!$A$2:$A${cattle_last}"


def cattle_formula(col_letter):
    idx = f"XLOOKUP($C$6,{CATTLE_RANGE},Cattle_by_State!{col_letter}$2:{col_letter}${cattle_last})"
    return f'=IF($C$6="","",IFERROR(IF({idx}=0,"n/a",{idx}),"n/a"))'


section("B10", "  2 · PRICES & MARGIN FOR THIS LOCATION")
out_label("B11", "\U0001F33E  Avg hay price")
out_value("C11",
          f'=IF($C$6="","",IFERROR(XLOOKUP($C$6,Hay_by_State!$A$2:$A${hay_last},'
          f'Hay_by_State!$B$2:$B${hay_last}),"n/a"))', '"$"#,##0" /ton"')
out_label("B12", "\U0001F404  Calf margin (region)")
out_value("C12",
          f'=IF($C$8="","",IFERROR(XLOOKUP($C$8,CalfMargin_by_Region!$A$2:$A${margin_last},'
          f'CalfMargin_by_Region!$B$2:$B${margin_last}),"n/a"))', '"$"#,##0.00" /cow"')

lk["B14"].value = "\U0001F4B2  Auction cattle averages (state)"
lk["B14"].font = Font(bold=True, size=11, color=NAVY)
r = 15
for j, lab in enumerate(CAT_LABELS, start=2):
    unit = "$/head" if "head" in lab else "$/cwt"
    out_label(f"B{r}", lab)
    out_value(f"C{r}", cattle_formula(get_column_letter(j)))
    uc = lk[f"D{r}"]
    uc.value, uc.font, uc.alignment = unit, Font(size=9, color="808080"), LEFT
    r += 1

fsec = r + 1
section(f"B{fsec}", "  3 · FENCING ESTIMATE")
field_label(f"B{fsec+1}", "Fence type:")
input_cell(f"C{fsec+1}", list(fence)[0])
field_label(f"B{fsec+2}", "Amount needed:")
input_cell(f"C{fsec+2}", 0)
field_label(f"B{fsec+3}", "Unit:")
input_cell(f"C{fsec+3}", "Miles")
amt, unit_c, ft = f"$C${fsec+2}", f"$C${fsec+3}", f"$C${fsec+1}"
lk[f"B{fsec+4}"].value = "Fence length (feet):"
lk[f"B{fsec+4}"].font, lk[f"B{fsec+4}"].alignment = Font(italic=True, size=10, color="595959"), RIGHT
lk[f"C{fsec+4}"] = (f'=IF({unit_c}="Feet",{amt},IF({unit_c}="Miles",{amt}*5280,'
                    f'IF({unit_c}="Acres",4*SQRT({amt}*43560),0)))')
lk[f"C{fsec+4}"].font, lk[f"C{fsec+4}"].alignment = Font(italic=True, size=10, color="595959"), CEN
lk[f"C{fsec+4}"].number_format = '#,##0" ft"'
out_label(f"B{fsec+5}", "Cost per foot")
out_value(f"C{fsec+5}",
          f'=IFERROR(XLOOKUP({ft},Fencing_Rates!$A$2:$A${fence_last},'
          f'Fencing_Rates!$B$2:$B${fence_last}),"n/a")', '"$"#,##0.00')
tc = lk[f"B{fsec+6}"]
tc.value, tc.font = "\U0001F4B0  Estimated total fencing cost", Font(bold=True, size=12, color=NAVY)
tc.fill, tc.alignment, tc.border = PatternFill("solid", fgColor=GREEN), LEFT, BORDER
res = lk[f"C{fsec+6}"]
res.value = f'=IFERROR(C{fsec+4}*C{fsec+5},"—")'
res.font = Font(bold=True, size=16, color=NAVY)
res.fill, res.alignment, res.border = PatternFill("solid", fgColor=GREEN), CEN, BORDER
res.number_format = '"$"#,##0'
lk[f"B{fsec+8}"].value = ("Note: for Acres, fence length assumes a square paddock "
                          "(perimeter = 4 × √(acres × 43,560 sq ft)). "
                          "Feet and Miles are exact.")
lk[f"B{fsec+8}"].font = Font(italic=True, size=9, color="808080")
lk.merge_cells(f"B{fsec+8}:F{fsec+8}")

dv_state = DataValidation(type="list", formula1="=StateList", allow_blank=True)
dv_county = DataValidation(type="list", formula1="=CountyList", allow_blank=True)
dv_fence = DataValidation(type="list", formula1=f"=Fencing_Rates!$A$2:$A${fence_last}", allow_blank=True)
dv_unit = DataValidation(type="list", formula1='"Acres,Miles,Feet"', allow_blank=True)
for dv in (dv_state, dv_county, dv_fence, dv_unit):
    lk.add_data_validation(dv)
dv_state.add(lk["C6"])
dv_county.add(lk["C7"])
dv_fence.add(lk[f"C{fsec+1}"])
dv_unit.add(lk[f"C{fsec+3}"])

for col, width in [("A", 2), ("B", 30), ("C", 20), ("D", 8), ("E", 4), ("F", 20)]:
    lk.column_dimensions[col].width = width

wb.save(OUT)
print(f"wrote {OUT} (cattle {len(cattle_states)}, hay {len(hay)}, margin {len(margin)}, "
      f"fence {len(fence)}, counties {len(counties)})")


if __name__ == "__main__":
    pass
