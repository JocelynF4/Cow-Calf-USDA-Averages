import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "averages_long.csv")
OUT = os.path.join(ROOT, "USDA_Cattle_2025_State_Averages.xlsx")

CATS = [
    ("Overall Feeder", "$/cwt"),
    ("Feeder Steers", "$/cwt"),
    ("Feeder Heifers", "$/cwt"),
    ("Feeder Bulls", "$/cwt"),
    ("Feeder Cows", "$/cwt"),
    ("Stock Cows ($/cwt)", "$/cwt"),
    ("Stock Cows ($/head)", "$/head"),
    ("Bred Cows ($/head)", "$/head"),
    ("Bred Cows ($/cwt)", "$/cwt"),
    ("Bred Heifers ($/head)", "$/head"),
    ("Bred Heifers ($/cwt)", "$/cwt"),
    ("Cow-Calf Pairs ($/head)", "$/head"),
    ("Heifer Pairs ($/head)", "$/head"),
    ("Replacement Bulls ($/head)", "$/head"),
    ("Replacement Bulls ($/cwt)", "$/cwt"),
]

NAVY, BLUE, LIGHT2, AMBER, GREY = "1F3A5F", "2F5C8F", "EAF1F8", "FFF2CC", "F2F2F2"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
SUB_FONT = Font(italic=True, size=10, color="595959")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")


def price_fmt(unit):
    return '"$"#,##0.00' if unit == "$/cwt" else '"$"#,##0'


def load():
    rows = list(csv.DictReader(open(SRC)))
    idx = {(r["state"], r["category"]): r for r in rows}
    states = sorted({r["state"] for r in rows})
    return rows, idx, states


def build_data_sheet(ws, idx, states):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "USDA Cattle Auction — 2025 Annual Average Prices by State"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = ("Head-count-weighted averages from USDA Market News weekly auction "
                "summaries, Jan–Dec 2025. $/cwt = per hundredweight; $/head = per "
                "animal. Blank = category not reported by that state.")
    ws["A2"].font = SUB_FONT
    ws.merge_cells("A1:P1")
    ws.merge_cells("A2:P2")

    hrow = 4
    ws.cell(hrow, 1, "State")
    for j, (label, _) in enumerate(CATS, start=2):
        ws.cell(hrow, j, label)
    wk_col = 2 + len(CATS)
    note_col = wk_col + 1
    ws.cell(hrow, wk_col, "Data weeks")
    ws.cell(hrow, note_col, "Coverage note")
    for col in range(1, note_col + 1):
        c = ws.cell(hrow, col)
        c.font, c.alignment, c.border = HDR_FONT, CENTER, BORDER
        c.fill = PatternFill("solid", fgColor=BLUE)

    first = hrow + 1
    for i, state in enumerate(states):
        r = first + i
        a = ws.cell(r, 1, state)
        a.font, a.alignment, a.border = Font(bold=True, size=10), LEFT, BORDER
        a.fill = PatternFill("solid", fgColor=GREY)
        weeks = []
        overall_weeks = None
        for j, (label, unit) in enumerate(CATS, start=2):
            cell = ws.cell(r, j)
            rec = idx.get((state, label))
            if rec:
                cell.value = float(rec["avg_price"])
                cell.number_format = price_fmt(unit)
                weeks.append(int(rec["n_weeks"]))
                if label == "Overall Feeder":
                    overall_weeks = int(rec["n_weeks"])
            cell.alignment, cell.border = CENTER, BORDER
            cell.fill = PatternFill("solid", fgColor=LIGHT2 if j % 2 == 0 else "FFFFFF")
        lo, hi = (min(weeks), max(weeks)) if weeks else (0, 0)
        wc = ws.cell(r, wk_col, f"{lo}–{hi}" if lo != hi else f"{lo}")
        wc.alignment, wc.border = CENTER, BORDER
        note = ""
        if overall_weeks is not None and overall_weeks < 45:
            note = f"Partial year (Overall = {overall_weeks} wks) — interpret with caution"
        nc = ws.cell(r, note_col, note)
        nc.alignment, nc.border = LEFT, BORDER
        if note:
            for col in (1, wk_col, note_col):
                ws.cell(r, col).fill = PatternFill("solid", fgColor=AMBER)

    ws.column_dimensions["A"].width = 16
    for j in range(2, 2 + len(CATS)):
        ws.column_dimensions[get_column_letter(j)].width = 13
    ws.column_dimensions[get_column_letter(wk_col)].width = 11
    ws.column_dimensions[get_column_letter(note_col)].width = 42
    ws.freeze_panes = ws.cell(first, 2)
    ws.row_dimensions[hrow].height = 44
    return hrow, first, first + len(states) - 1


def build_lookup_sheet(lk, states, hrow, first, last):
    lk.sheet_view.showGridLines = False
    lk["B2"] = "Cattle Price Lookup"
    lk["B2"].font = Font(bold=True, size=18, color=NAVY)
    lk["B3"] = "Pick a state and a category to see its 2025 head-weighted average."
    lk["B3"].font = SUB_FONT

    def label_cell(ref, txt):
        c = lk[ref]
        c.value, c.font = txt, Font(bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = BORDER

    def input_cell(ref, value):
        c = lk[ref]
        c.value = value
        c.fill = PatternFill("solid", fgColor=AMBER)
        c.alignment, c.border, c.font = CENTER, BORDER, Font(bold=True, size=11)

    label_cell("B5", "State:")
    input_cell("C5", states[0])
    label_cell("B6", "Category:")
    input_cell("C6", "Overall Feeder")
    lk.merge_cells("C5:D5")
    lk.merge_cells("C6:D6")

    last_col = get_column_letter(1 + len(CATS))
    dv_s = DataValidation(type="list", formula1=f"Data!$A${first}:$A${last}", allow_blank=False)
    dv_c = DataValidation(type="list", formula1=f"Data!$B${hrow}:${last_col}${hrow}", allow_blank=False)
    lk.add_data_validation(dv_s)
    lk.add_data_validation(dv_c)
    dv_s.add(lk["C5"])
    dv_c.add(lk["C6"])

    val_range = f"Data!$B${first}:${last_col}${last}"
    match_s = f"MATCH($C$5,Data!$A${first}:$A${last},0)"
    match_c = f"MATCH($C$6,Data!$B${hrow}:${last_col}${hrow},0)"
    index_expr = f"INDEX({val_range},{match_s},{match_c})"

    lk["B8"] = "Average price:"
    lk["B8"].font = Font(bold=True, size=12, color=NAVY)
    lk["B8"].alignment = Alignment(horizontal="right")
    res = lk["C8"]
    res.value = f'=IFERROR(IF({index_expr}=0,"— not reported —",{index_expr}),"— not reported —")'
    res.font = Font(bold=True, size=20, color=NAVY)
    res.fill = PatternFill("solid", fgColor="DCE6F1")
    res.alignment, res.border = CENTER, BORDER
    lk.merge_cells("C8:D8")
    lk["E8"] = '=IF(ISNUMBER(SEARCH("head",$C$6)),"$/head",IF($C$6="","","$/cwt"))'
    lk["E8"].font = Font(bold=True, size=12, color="595959")
    lk["E8"].alignment = LEFT

    sd = "'Source Detail'"
    weeks_sumifs = f"SUMIFS({sd}!$G:$G,{sd}!$A:$A,$C$5,{sd}!$B:$B,$C$6)"
    head_sumifs = f"SUMIFS({sd}!$E:$E,{sd}!$A:$A,$C$5,{sd}!$B:$B,$C$6)"
    lk["B10"] = "Weeks of data behind it:"
    lk["B10"].font = Font(size=10, color="595959")
    lk["B10"].alignment = Alignment(horizontal="right")
    lk["C10"] = f'=IFERROR(IF({weeks_sumifs}=0,"",{weeks_sumifs}),"")'
    lk["C10"].font = Font(size=10, color="595959")
    lk["C10"].alignment = CENTER
    lk["B11"] = "Head sold behind it:"
    lk["B11"].font = Font(size=10, color="595959")
    lk["B11"].alignment = Alignment(horizontal="right")
    lk["C11"] = f'=IFERROR(IF({head_sumifs}=0,"",TEXT({head_sumifs},"#,##0")),"")'
    lk["C11"].font = Font(size=10, color="595959")
    lk["C11"].alignment = CENTER
    lk["B13"] = ('=IF($C$10="","",IF($C$10<45,"⚠ Partial-year data — this state did '
                 'not report all 52 weeks; average may be seasonally skewed.",""))')
    lk["B13"].font = Font(italic=True, size=10, color="C00000")
    lk.merge_cells("B13:H13")
    for col, width in [("A", 3), ("B", 22), ("C", 16), ("D", 10), ("E", 10),
                       ("F", 6), ("G", 6), ("H", 10)]:
        lk.column_dimensions[col].width = width
    lk.row_dimensions[8].height = 32


def build_detail_sheet(dt, rows):
    dt.sheet_view.showGridLines = False
    heads = ["State", "Category", "Unit", "Avg Price", "Total Head",
             "# Observations", "# Weeks"]
    for j, h in enumerate(heads, start=1):
        c = dt.cell(1, j, h)
        c.font, c.alignment, c.border = HDR_FONT, CENTER, BORDER
        c.fill = PatternFill("solid", fgColor=BLUE)
    order = {label: i for i, (label, _) in enumerate(CATS)}
    r = 2
    for row in sorted(rows, key=lambda x: (x["state"], order.get(x["category"], 99))):
        vals = [row["state"], row["category"], row["unit"], float(row["avg_price"]),
                int(row["total_head"]), int(row["n_obs"]), int(row["n_weeks"])]
        for j, v in enumerate(vals, start=1):
            c = dt.cell(r, j, v)
            c.border = BORDER
            c.alignment = CENTER if j > 2 else LEFT
            if j == 4:
                c.number_format = price_fmt(row["unit"])
            if j == 5:
                c.number_format = "#,##0"
        r += 1
    for col, width in [("A", 16), ("B", 26), ("C", 9), ("D", 12), ("E", 13),
                       ("F", 15), ("G", 9)]:
        dt.column_dimensions[col].width = width
    dt.freeze_panes = "A2"
    dt.auto_filter.ref = f"A1:G{r - 1}"


def main():
    rows, idx, states = load()
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    hrow, first, last = build_data_sheet(data, idx, states)
    build_lookup_sheet(wb.create_sheet("Lookup"), states, hrow, first, last)
    build_detail_sheet(wb.create_sheet("Source Detail"), rows)
    wb.save(OUT)
    print(f"wrote {OUT} ({len(states)} states, {len(CATS)} categories, {len(rows)} detail rows)")


if __name__ == "__main__":
    main()
